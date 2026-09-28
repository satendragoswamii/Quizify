"""Render a :class:`ParseResult` as Excel, CSV, or JSON."""

import csv
import io
import json
from typing import Any, Dict, List, Tuple

from .models import ParseResult, Question

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    XLSX_OK = True
except ImportError:  # pragma: no cover
    XLSX_OK = False


HEADER_FILL = "FF1F3864"
LOW_CONFIDENCE_FILL = "FFFFF2CC"
MISSING_ANSWER_FILL = "FFFCE4E4"

BASE_COLUMNS = ["#", "Subject", "Topic", "Type", "Question"]
TAIL_COLUMNS = ["Answer", "Answer Text", "Explanation", "Marks", "Confidence", "Review Notes"]

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# How the Answer column renders a correct-option label.
#   "label"  -> the letter itself, e.g. "A" or "A,C"  (default, original behaviour)
#   "answer" -> the option's position as "Answer1", "Answer2", ...
#   "choice" -> the option's position as "Choice 1", "Choice 2", ...
ANSWER_STYLES = ("label", "answer", "choice")
_STYLE_TEMPLATES = {"answer": "Answer{}", "choice": "Choice {}"}


def _normalise_answer_style(answer_style: str) -> str:
    style = (answer_style or "label").lower()
    return style if style in ANSWER_STYLES else "label"


def _label_position(label: str) -> int:
    """Turn an A-Z option label into its 1-based position (A->1, B->2, ...).

    Returns 0 for anything that is not a single letter, so the caller can fall
    back to the raw label rather than emit a meaningless "Answer0".
    """
    label = (label or "").strip().upper()
    return ord(label) - 64 if len(label) == 1 and "A" <= label <= "Z" else 0


def _format_answer(question: Question, answer_style: str) -> str:
    """Render the Answer cell in the requested style.

    Only letter labels (A, B, C, ...) are remapped to positions; a text answer,
    or a label the app never assigned a letter to, is passed through unchanged.
    """
    template = _STYLE_TEMPLATES.get(answer_style)
    if template is None or not question.answer_labels:
        return question.answer_display
    parts = []
    for label in question.answer_labels:
        position = _label_position(label)
        parts.append(template.format(position) if position else label)
    return ",".join(parts)


def _column_names(max_options: int) -> List[str]:
    options = [f"Option {chr(64 + i)}" for i in range(1, max_options + 1)]
    return BASE_COLUMNS + options + TAIL_COLUMNS


def _row(question: Question, subject: str, topic: str, max_options: int,
         answer_style: str) -> List[Any]:
    row: List[Any] = [
        question.number,
        subject,
        topic,
        question.qtype.value,
        question.text,
    ]
    for i in range(max_options):
        row.append(question.options[i].text if i < len(question.options) else "")
    row.extend([
        _format_answer(question, answer_style),
        question.answer_text or "",
        question.explanation or "",
        question.marks if question.marks is not None else "",
        round(question.confidence, 2),
        "; ".join(question.warnings),
    ])
    return row


def build_rows(result: ParseResult, subject: str, topic: str, max_options: int,
               answer_style: str = "label"
               ) -> Tuple[List[str], List[List[Any]]]:
    max_options = max(2, min(int(max_options), 26))
    answer_style = _normalise_answer_style(answer_style)
    headers = _column_names(max_options)
    rows = [_row(q, subject, topic, max_options, answer_style) for q in result.questions]
    return headers, rows


# --------------------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------------------

def _style_sheet(ws, headers: List[str], rows: List[List[Any]]) -> None:
    header_font = Font(bold=True, color="FFFFFFFF", size=11)
    fill = PatternFill("solid", fgColor=HEADER_FILL)
    for col, name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {"#": 6, "Subject": 16, "Topic": 18, "Type": 14, "Question": 60,
              "Answer": 10, "Answer Text": 20, "Explanation": 40, "Marks": 8,
              "Confidence": 11, "Review Notes": 34}
    for col, name in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(name, 28)

    ws.freeze_panes = "F2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def _highlight_review_rows(ws, headers: List[str], questions: List[Question]) -> None:
    """Tint rows that a human should check: no answer, or low parser confidence."""
    warn = PatternFill("solid", fgColor=LOW_CONFIDENCE_FILL)
    missing = PatternFill("solid", fgColor=MISSING_ANSWER_FILL)
    for offset, question in enumerate(questions, 2):
        if not question.answer_display:
            style = missing
        elif question.confidence < 0.6 or question.warnings:
            style = warn
        else:
            continue
        for col in range(1, len(headers) + 1):
            ws.cell(row=offset, column=col).fill = style


def _summary_sheet(wb, result: ParseResult, subject: str, topic: str) -> None:
    ws = wb.create_sheet("Summary")
    stats = result.stats
    bold = Font(bold=True)

    entries: List[Tuple[str, Any]] = [
        ("Source", result.source_name or "pasted text"),
        ("Source format", stats["source_format"] or "text"),
        ("Parsed by", stats["engine"]),
        ("Subject", subject),
        ("Topic", topic),
        ("", ""),
        ("Total questions", stats["total"]),
        ("With answers", stats["answered"]),
        ("Missing answers", stats["unanswered"]),
        ("Needs review", stats["low_confidence"]),
        ("Average confidence", stats["avg_confidence"]),
        ("", ""),
    ]
    entries.extend(("Type: " + name, count) for name, count in sorted(stats["by_type"].items()))
    if result.warnings:
        entries.append(("", ""))
        entries.extend(("Warning", w) for w in result.warnings)

    for r, (label, value) in enumerate(entries, 1):
        ws.cell(row=r, column=1, value=label).font = bold
        ws.cell(row=r, column=2, value=value)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60


def to_excel(result: ParseResult, subject: str, topic: str, max_options: int,
             answer_style: str = "label") -> io.BytesIO:
    if not XLSX_OK:
        raise RuntimeError("Excel export needs the 'openpyxl' package.")
    headers, rows = build_rows(result, subject, topic, max_options, answer_style)

    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    _style_sheet(ws, headers, rows)
    _highlight_review_rows(ws, headers, result.questions)
    _summary_sheet(wb, result, subject, topic)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# --------------------------------------------------------------------------------------
# CSV / JSON
# --------------------------------------------------------------------------------------

def to_csv(result: ParseResult, subject: str, topic: str, max_options: int,
           answer_style: str = "label") -> io.BytesIO:
    headers, rows = build_rows(result, subject, topic, max_options, answer_style)
    text = io.StringIO()
    writer = csv.writer(text, quoting=csv.QUOTE_MINIMAL, lineterminator="\n")
    writer.writerow(headers)
    writer.writerows(rows)
    # BOM so Excel opens UTF-8 CSVs with the right encoding.
    return io.BytesIO(text.getvalue().encode("utf-8-sig"))


def to_json(result: ParseResult, subject: str, topic: str, max_options: int,
            answer_style: str = "label") -> io.BytesIO:
    payload: Dict[str, Any] = result.to_dict()
    payload["subject"] = subject
    payload["topic"] = topic
    answer_style = _normalise_answer_style(answer_style)
    if answer_style != "label":
        payload["answer_style"] = answer_style
        for question, rendered in zip(payload["questions"], result.questions):
            question["answer"] = _format_answer(rendered, answer_style)
    return io.BytesIO(json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8"))


_EXPORTERS = {
    "excel": (to_excel, MIME_XLSX, "xlsx"),
    "xlsx": (to_excel, MIME_XLSX, "xlsx"),
    "csv": (to_csv, "text/csv; charset=utf-8", "csv"),
    "json": (to_json, "application/json", "json"),
}


def export(result: ParseResult, output_format: str, subject: str, topic: str,
           max_options: int = 4, basename: str = "questions",
           answer_style: str = "label") -> Tuple[io.BytesIO, str, str]:
    """Render ``result``. Returns ``(buffer, mimetype, filename)``.

    ``answer_style`` controls how the Answer column renders a correct option:
    ``"label"`` keeps the letter (A, B, ...), ``"answer"`` writes ``Answer1``,
    ``"choice"`` writes ``Choice 1`` — position taken from the option's letter.
    """
    exporter, mimetype, extension = _EXPORTERS.get(
        (output_format or "excel").lower(), _EXPORTERS["excel"]
    )
    buffer = exporter(result, subject, topic, max_options, answer_style)
    return buffer, mimetype, f"{basename}.{extension}"
