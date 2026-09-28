"""Quizify test suite.

Run with:  python tests/test_quizify.py

Covers the format matrix (how a quiz can be written), the extraction matrix (what
file it arrives in), the export layer, and the HTTP endpoints. The AI assist is
switched off throughout so these tests never make a network call.
"""

import csv as csvmod
import io
import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Use a throwaway admin database so these tests never touch the real one.
os.environ.setdefault(
    "QUIZ_DB_PATH",
    os.path.join(tempfile.mkdtemp(prefix="quizify-parse-test-"), "test.db"))

import quizify  # noqa: E402

PASSED, FAILED = [], []


def _describe(result):
    return " | ".join(
        f"[{q.qtype.value}] {q.text[:45]!r} opts={[o.text[:14] for o in q.options]} "
        f"ans={q.answer_display!r}"
        for q in result.questions[:6]
    )


def check(name, text=None, *, data=None, filename=None, questions=None, answers=None,
          types=None, option_counts=None):
    """Parse an input and assert on the shape of the result."""
    try:
        if data is not None:
            result = quizify.parse_file(data, filename, ai_mode="off")
        else:
            result = quizify.parse_text(text, ai_mode="off")
    except Exception as exc:
        FAILED.append(f"{name}: raised {type(exc).__name__}: {exc}")
        return None

    problems = []
    if questions is not None and len(result.questions) != questions:
        problems.append(f"question count {len(result.questions)} != {questions}")
    if answers is not None:
        got = [q.answer_display for q in result.questions]
        if got != answers:
            problems.append(f"answers {got} != {answers}")
    if types is not None:
        got = [q.qtype.value for q in result.questions]
        if got != types:
            problems.append(f"types {got} != {types}")
    if option_counts is not None:
        got = [len(q.options) for q in result.questions]
        if got != option_counts:
            problems.append(f"option counts {got} != {option_counts}")

    if problems:
        FAILED.append(f"{name}: {'; '.join(problems)}\n      got: {_describe(result)}")
    else:
        PASSED.append(name)
    return result


def expect(name, condition, detail=""):
    (PASSED if condition else FAILED).append(name if condition else f"{name}: {detail}")


# --------------------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------------------

def docx_with_bold_answers():
    from docx import Document
    doc = Document()
    doc.add_paragraph("Q1. Which is a fruit?")
    doc.add_paragraph("A. Carrot")
    doc.add_paragraph().add_run("B. Apple").bold = True
    doc.add_paragraph("C. Potato")
    doc.add_paragraph("D. Onion")
    doc.add_paragraph("")
    doc.add_paragraph("Q2. Which is a metal?")
    doc.add_paragraph("A. Oxygen")
    doc.add_paragraph("B. Nitrogen")
    doc.add_paragraph().add_run("C. Iron").bold = True
    doc.add_paragraph("D. Helium")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def docx_all_bold():
    """Whole document bold — bold must not be read as an answer cue."""
    from docx import Document
    doc = Document()
    for line in ["Q1. Which is a fruit?", "A. Carrot", "B. Apple", "C. Potato", "D. Onion"]:
        doc.add_paragraph().add_run(line).bold = True
    doc.add_paragraph("Answer: B")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def docx_duplicated_in_table():
    from docx import Document
    doc = Document()
    doc.add_paragraph("Q1. What is 1+1?")
    doc.add_paragraph("A. 1")
    doc.add_paragraph("B. 2")
    doc.add_paragraph("Answer: B")
    table = doc.add_table(rows=4, cols=1)
    for row, value in enumerate(["Q1. What is 1+1?", "A. 1", "B. 2", "Answer: B"]):
        table.cell(row, 0).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def xlsx_structured():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Question", "Option A", "Option B", "Option C", "Answer"])
    ws.append(["What is 5*5?", "20", "25", "30", "B"])
    ws.append(["What is 9-4?", "5", "6", "7", "A"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def minimal_pdf(lines):
    """A single-page PDF built by hand, so the tests need no PDF-writing dependency."""
    content = "BT /F1 11 Tf 72 730 Td 15 TL\n"
    for line in lines:
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content += f"({escaped}) Tj T*\n"
    content += "ET"
    stream = content.encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF").encode()
    return bytes(out)


# --------------------------------------------------------------------------------------
# Question / option / answer formats
# --------------------------------------------------------------------------------------

check("Q1. numbering with A. options and Answer: line", """
Q1. Which of the following is a programming language?
A. HTML
B. Python
C. CSS
D. XML
Answer: B

Q2. What does API stand for?
A. Application Programming Interface
B. Applied Program Integration
C. Automatic Protocol Interface
D. Advanced Programming Instruction
Answer: A
""", questions=2, answers=["B", "A"], types=["MCQ", "MCQ"], option_counts=[4, 4])

check("numbered questions with (a) options", """
1. The capital of France is
(a) Berlin
(b) Madrid
(c) Paris
(d) Rome
Ans: c

2. Water boils at
(a) 50 C
(b) 100 C
(c) 150 C
(d) 200 C
Ans: b
""", questions=2, answers=["C", "B"], option_counts=[4, 4])

check("numbered questions and numbered options together", """
1. Which is largest?
1) Mercury
2) Venus
3) Jupiter
4) Mars
Answer: 3

2. Which is smallest?
1) Mercury
2) Venus
3) Jupiter
4) Mars
Answer: 1
""", questions=2, answers=["C", "A"], option_counts=[4, 4])

check("alternative question prefixes", """
Que. 1 What is HTML?
A. A markup language
B. A programming language
C. A database
D. An OS
Answer: A

Q.No.2 What is CSS?
A. A database
B. A stylesheet language
C. A compiler
D. A browser
Answer: B
""", questions=2, answers=["A", "B"])

check("bracket option labels", """
Q1. Pick the odd one out.
[A] Cat
[B] Dog
[C] Table
[D] Horse
Answer: C
""", questions=1, option_counts=[4], answers=["C"])

check("roman numeral options", """
Q1. Which is the smallest?
i. 100
ii. 10
iii. 1
iv. 50
Answer: iii
""", questions=1, option_counts=[4], answers=["C"])

check("26 options", "Q1. Pick one.\n"
      + "\n".join(f"{chr(65 + i)}. Choice {i + 1}" for i in range(26))
      + "\nAnswer: Z", questions=1, option_counts=[26], answers=["Z"])

check("options with no labels at all", """
Which planet is closest to the sun?
Mercury
Venus
Earth
Mars

What is the largest mammal?
Elephant
Blue whale
Giraffe
Hippo
""", questions=2, option_counts=[4, 4])

check("markdown bullets", """
### Quiz

**Q1.** Which language runs in a browser?
- A. Python
- B. JavaScript
- C. C++
- D. Java

Answer: B
""", questions=1, answers=["B"])

check("options wrapped over several lines", """
Q1. Which statement best describes photosynthesis?
A. The process by which green plants and some other organisms
   use sunlight to synthesize foods from carbon dioxide
   and water, generating oxygen as a byproduct
B. The process of cellular respiration
C. The process of digestion
D. The process of evaporation
Answer: A
""", questions=1, option_counts=[4], answers=["A"])


# --------------------------------------------------------------------------------------
# Answer sources
# --------------------------------------------------------------------------------------

check("trailing answer key section", """
Q1. Largest planet?
A. Earth
B. Jupiter
C. Mars
D. Venus

Q2. Smallest prime?
A. 0
B. 1
C. 2
D. 3

Q3. Chemical symbol for gold?
A. Ag
B. Au
C. Gd
D. Go

Answer Key
Q1: b, Q2: c, Q3: b
""", questions=3, answers=["B", "C", "B"])

check("dash-style answer grid", """
1. Fastest land animal?
A. Lion
B. Cheetah
C. Horse
D. Tiger

2. Largest ocean?
A. Atlantic
B. Indian
C. Arctic
D. Pacific

Answers
1-B, 2-D
""", questions=2, answers=["B", "D"])

check("answer key listed out of order", """
Q1. Capital of India?
A. Mumbai
B. Delhi
C. Chennai
D. Kolkata

Q2. Capital of Australia?
A. Sydney
B. Melbourne
C. Canberra
D. Perth

ANSWER KEY
Q2: C
Q1: B
""", questions=2, answers=["B", "C"])

check("asterisk marks the correct option", """
Q1. Which is a fruit?
A. Carrot
B. Apple *
C. Potato
D. Onion

Q2. Capital of Italy?
A. Rome *
B. Milan
C. Naples
D. Turin
""", questions=2, answers=["B", "A"])

check("bold marks the correct option (docx)", data=docx_with_bold_answers(),
      filename="quiz.docx", questions=2, answers=["B", "C"], option_counts=[4, 4])

check("uniformly bold document ignores the bold cue", data=docx_all_bold(),
      filename="quiz.docx", questions=1, answers=["B"])

check("True/False answer key", """
Q1. Water boils at 100 C. (True/False)

Q2. The sun orbits the Earth. (True/False)

Answer Key
Q1: True, Q2: False
""", questions=2, answers=["A", "B"], types=["True/False", "True/False"])

check("numbered Yes/No options are not an answer grid", """
Q1. Is the sky blue?
1. Yes
2. No
Answer: 1

Q2. Is fire cold?
1. Yes
2. No
Answer: 2
""", questions=2, answers=["A", "B"], option_counts=[2, 2])


# --------------------------------------------------------------------------------------
# Question types
# --------------------------------------------------------------------------------------

check("True/False from the stem", """
Q1. The Earth is flat. (True/False)
Answer: False

Q2. Python is a programming language. (True/False)
Answer: True
""", questions=2, answers=["B", "A"], types=["True/False", "True/False"])

check("multi-select", """
Q1. Select all that apply. Which are prime numbers?
A. 2
B. 4
C. 5
D. 9
Answer: A, C
""", questions=1, answers=["A,C"], types=["Multi-Select"])

check("fill in the blank", """
Q1. The capital of Japan is ________.
Answer: Tokyo

Q2. Water freezes at _____ degrees Celsius.
Answer: 0
""", questions=2, types=["Fill in the Blank", "Fill in the Blank"], answers=["Tokyo", "0"])

check("assertion and reason", """
Q1. Assertion (A): Water boils at 100 C at sea level. Reason (R): Atmospheric pressure is 1 atm at sea level.
A. Both A and R are true and R explains A
B. Both A and R are true but R does not explain A
C. A is true, R is false
D. A is false, R is true
Answer: A
""", questions=1, types=["Assertion-Reason"], answers=["A"])

check("matching", """
Q1. Match the following. Column A with Column B.
A. Newton - Gravity
B. Einstein - Relativity
C. Darwin - Evolution
D. Curie - Radioactivity
Answer: A
""", questions=1, types=["Matching"])

check("ordering", """
Q1. Arrange the following in the correct chronological order.
A. 1947, 1950, 1971
B. 1950, 1947, 1971
C. 1971, 1950, 1947
D. 1947, 1971, 1950
Answer: A
""", questions=1, types=["Ordering"], answers=["A"])

check("short answer", """
Q1. Define osmosis.
Correct Answer: Movement of solvent across a semipermeable membrane.

Q2. Who wrote Hamlet?
Answer: William Shakespeare
""", questions=2, types=["Short Answer", "Short Answer"])

check("several types in one document", """
Q1. Python is compiled. (True/False)
Answer: False

Q2. The capital of France is ______.
Answer: Paris

Q3. Which are mammals? Select all that apply.
A. Whale
B. Shark
C. Bat
D. Trout
Answer: A, C

Q4. Define entropy.
Answer: A measure of disorder in a system.

Q5. What is 12 * 12?
A. 124
B. 144
C. 154
D. 164
Answer: B
""", questions=5,
   types=["True/False", "Fill in the Blank", "Multi-Select", "Short Answer", "MCQ"])


# --------------------------------------------------------------------------------------
# File formats
# --------------------------------------------------------------------------------------

check("csv with question/option columns", data=(
    b"Question,Option A,Option B,Option C,Option D,Answer,Explanation\n"
    b"What is 2+2?,3,4,5,6,B,Simple addition\n"
    b"Capital of Japan?,Osaka,Kyoto,Tokyo,Nagoya,C,It moved in 1868\n"
), filename="quiz.csv", questions=2, answers=["B", "C"], option_counts=[4, 4])

check("xlsx with question/option columns", data=xlsx_structured(), filename="quiz.xlsx",
      questions=2, answers=["B", "A"])

check("structured json", data=json.dumps({"questions": [
    {"question": "What is 3+3?", "options": ["5", "6", "7"], "answer": "B"},
    {"question": "Largest planet?", "options": ["Earth", "Jupiter"], "answer": 2},
]}).encode(), filename="quiz.json", questions=2, answers=["B", "B"])

check("html", data=(
    b"<html><body><p>Q1. What is the boiling point of water?</p>"
    b"<ul><li>A. 50 C</li><li>B. 100 C</li><li>C. 150 C</li><li>D. 200 C</li></ul>"
    b"<p>Answer: B</p></body></html>"
), filename="quiz.html", questions=1, answers=["B"])

check("pdf", data=minimal_pdf([
    "Sample Quiz - Page 1 of 1", "",
    "Q1. What is the capital of Kenya?",
    "A. Lagos", "B. Nairobi", "C. Kampala", "D. Addis Ababa", "Answer: B", "",
    "Q2. Which of these is a noble gas?",
    "A. Oxygen", "B. Nitrogen", "C. Neon", "D. Hydrogen", "",
    "Answer Key", "Q2: c",
]), filename="quiz.pdf", questions=2, answers=["B", "C"], option_counts=[4, 4])

check("file contents win over a wrong extension", data=docx_with_bold_answers(),
      filename="quiz.txt", questions=2, answers=["B", "C"])

check("duplicates across paragraphs and tables are merged",
      data=docx_duplicated_in_table(), filename="quiz.docx", questions=1)


# --------------------------------------------------------------------------------------
# Robustness
# --------------------------------------------------------------------------------------

check("page furniture is discarded", """
Page 1 of 3
=====================

Q1. Capital of Spain?
A. Lisbon
B. Madrid
C. Paris
D. Rome
Answer: B

Page 2 of 3
-----------

Q2. Capital of Portugal?
A. Lisbon
B. Madrid
C. Paris
D. Rome
Answer: A
""", questions=2, answers=["B", "A"])

check("years in prose are not question numbers", """
Q1. In which year did World War II end?
A. 1943
B. 1945
C. 1947
D. 1950
Answer: B

Q2. 1969 was the year of the moon landing. Which agency achieved it?
A. ESA
B. NASA
C. ISRO
D. JAXA
Answer: B
""", questions=2, answers=["B", "B"])

check("numbering that restarts per section", """
Section A
1. What is H2O?
A. Water
B. Salt
Answer: A

Section B
1. What is NaCl?
A. Water
B. Salt
Answer: B
""", questions=2, answers=["A", "B"])

result = check("decimals in option text", """
Q1. What is the value of pi to 2 decimal places?
A. 3.14
B. 3.41
C. 2.71
D. 1.62
Answer: A
""", questions=1, option_counts=[4], answers=["A"])
expect("decimal option text preserved verbatim",
       bool(result and result.questions[0].options[0].text == "3.14"))

result = check("non-English content", """
Q1. भारत की राजधानी क्या है?
A. मुंबई
B. दिल्ली
C. चेन्नई
D. कोलकाता
Answer: B

Q2. ¿Cuál es la capital de España?
A. Barcelona
B. Madrid
C. Sevilla
D. Valencia
Answer: B
""", questions=2, answers=["B", "B"])
expect("non-Latin option text preserved",
       bool(result and result.questions[0].options[1].text == "दिल्ली"))

result = check("smart quotes and em dashes", """
Q1. Who said “I think, therefore I am”?
A. Plato
B. Descartes — the French philosopher
C. Kant
D. Hume
Answer: B
""", questions=1, option_counts=[4])
expect("smart quotes normalised", bool(result and '"I think' in result.questions[0].text))

result = check("explanation lines", """
Q1. What is 2+2?
A. 3
B. 4
C. 5
D. 6
Answer: B
Explanation: Basic arithmetic addition of two and two.
""", questions=1, answers=["B"])
expect("explanation captured",
       bool(result and result.questions[0].explanation
            and "arithmetic" in result.questions[0].explanation))

result = check("marks annotation", """
Q1. Explain photosynthesis. (5 marks)
A. Plants make food
B. Plants eat soil
C. Plants sleep
D. Plants swim
Answer: A
""", questions=1, answers=["A"])
expect("marks captured", bool(result and result.questions[0].marks == 5.0),
       f"got {result.questions[0].marks if result else None}")

result = check("questions with no answers anywhere", """
Q1. What is gravity?
A. A force
B. A color
C. A sound
D. A taste

Q2. What is light?
A. A wave
B. A rock
C. A tree
D. A fish
""", questions=2, answers=["", ""])
expect("missing answers are flagged",
       bool(result and all("No answer found." in q.warnings for q in result.questions)))

for label, junk in [("empty", ""), ("whitespace", "   \n\n  \n"),
                    ("prose only", "The quick brown fox jumps over the lazy dog.")]:
    result = quizify.parse_text(junk, ai_mode="off")
    expect(f"junk input handled: {label}",
           not result.questions and bool(result.warnings),
           f"got {len(result.questions)} questions")


# --------------------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------------------

export_source = quizify.parse_text("""
Q1. Capital of Japan?
A. Osaka
B. Tokyo
C. Kyoto
Answer: B
Explanation: Tokyo has been the capital since 1868.

Q2. Unanswered question here?
A. Yes
B. No
""", ai_mode="off")

try:
    import openpyxl
    buffer, mimetype, filename = quizify.export(
        export_source, "excel", "Geography", "Capitals", max_options=5)
    wb = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    ws = wb["Questions"]
    headers = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    expect("excel: summary sheet present", "Summary" in wb.sheetnames)
    expect("excel: header layout", headers[:5] == ["#", "Subject", "Topic", "Type", "Question"])
    expect("excel: max_options honoured", "Option E" in headers)
    expect("excel: review columns present",
           "Confidence" in headers and "Review Notes" in headers)
    expect("excel: subject and topic written", row[1] == "Geography" and row[2] == "Capitals")
    expect("excel: answer written", row[headers.index("Answer")] == "B")
    expect("excel: explanation written", "1868" in str(row[headers.index("Explanation")]))
    expect("excel: panes frozen", ws.freeze_panes == "F2")
    expect("excel: unanswered row highlighted",
           ws.cell(row=3, column=1).fill.fgColor.rgb == "FFFCE4E4")
    expect("excel: filename", filename.endswith(".xlsx"))
except Exception as exc:
    FAILED.append(f"excel export: raised {type(exc).__name__}: {exc}")

try:
    buffer, mimetype, filename = quizify.export(export_source, "csv", "S", "T", max_options=3)
    rows = list(csvmod.reader(io.StringIO(buffer.getvalue().decode("utf-8-sig"))))
    expect("csv: header and row count", rows[0][0] == "#" and len(rows) == 3)
    expect("csv: filename and mimetype", filename.endswith(".csv") and "csv" in mimetype)
except Exception as exc:
    FAILED.append(f"csv export: raised {type(exc).__name__}: {exc}")

try:
    buffer, mimetype, filename = quizify.export(export_source, "json", "S", "T")
    payload = json.loads(buffer.getvalue().decode())
    expect("json: stats and answers",
           payload["stats"]["total"] == 2
           and payload["questions"][0]["answer"] == "B"
           and payload["subject"] == "S")
except Exception as exc:
    FAILED.append(f"json export: raised {type(exc).__name__}: {exc}")


# --------------------------------------------------------------------------------------
# HTTP endpoints
# --------------------------------------------------------------------------------------

os.environ.pop("ANTHROPIC_API_KEY", None)
os.environ["QUIZ_API_ENABLED"] = "false"

import app as flask_app  # noqa: E402

flask_app.app.config["TESTING"] = True
client = flask_app.app.test_client()

# The app now requires a signed-in user for the converter and its APIs. These endpoint
# tests exercise staff functionality, so authenticate the client as an admin. (A student
# would be sandboxed away from the converter, which is not what these tests check.)
from quizify.admin import auth as _auth, repo as _repo  # noqa: E402
with flask_app.app.app_context():
    _staff = _repo.get_user_by_email("tester@quizify.local")
    if _staff is None:
        _staff_id = _repo.create_user("tester@quizify.local", "tester-pass-123",
                                      name="Test Admin", role="admin")
    else:
        _staff_id = _staff["id"]
with client.session_transaction() as _sess:
    _sess[_auth.SESSION_USER_KEY] = _staff_id

SAMPLE = "Q1. Capital of Peru?\nA. Lima\nB. Quito\nC. La Paz\nAnswer: A"

response = client.get("/api/health")
expect("GET /api/health",
       response.status_code == 200 and "docx" in response.get_json()["supported_formats"],
       f"status {response.status_code}")

response = client.post("/api/parse", data={
    "subject": "Geo", "topic": "SA", "quiz_text": SAMPLE, "ai_mode": "off"})
body = response.get_json() if response.status_code == 200 else {}
expect("POST /api/parse",
       response.status_code == 200 and body["stats"]["total"] == 1
       and body["questions"][0]["answer"] == "A",
       f"status {response.status_code}")

response = client.post("/", data={
    "subject": "Geo", "topic": "SA", "quiz_text": SAMPLE,
    "format": "excel", "max_options": "4"})
disposition = response.headers.get("Content-Disposition", "")
expect("POST / returns an xlsx download",
       response.status_code == 200 and "attachment" in disposition
       and ".xlsx" in disposition,
       f"status {response.status_code} disposition {disposition}")

response = client.post("/", data={"subject": "S", "topic": "T", "format": "csv",
                                  "max_options": "4"})
expect("POST / with no input returns 400",
       response.status_code == 400 and b"upload a file or paste" in response.data,
       f"status {response.status_code}")

response = client.post("/", data={
    "subject": "S", "topic": "T", "format": "excel", "max_options": "4",
    "file": (io.BytesIO(b"binary junk"), "notes.exe")})
expect("POST / rejects unsupported extensions",
       response.status_code == 400 and b"not accepted" in response.data,
       f"status {response.status_code}")

response = client.post("/", data={
    "subject": "S", "topic": "T", "format": "excel", "max_options": "999",
    "quiz_text": SAMPLE})
expect("max_options out of range is clamped", response.status_code == 200,
       f"status {response.status_code}")

response = client.post("/api/parse", data={
    "subject": "S", "topic": "T", "ai_mode": "off",
    "file": (io.BytesIO(b"%PDF-1.4 not really a pdf"), "x.pdf")})
expect("corrupt pdf returns a clean error",
       response.status_code == 400 and "error" in (response.get_json() or {}),
       f"status {response.status_code}")

response = client.post("/api/chat", json={"message": "How do I use this?"})
expect("chat degrades gracefully with no key",
       response.status_code == 503 and (response.get_json() or {}).get("fallback"),
       f"status {response.status_code}")


# --------------------------------------------------------------------------------------
# Report
# --------------------------------------------------------------------------------------

print(f"\n{'=' * 72}")
print(f"PASSED: {len(PASSED)}    FAILED: {len(FAILED)}")
print("=" * 72)
for name in PASSED:
    print(f"  ok    {name}")
if FAILED:
    print()
    for failure in FAILED:
        print(f"  FAIL  {failure}")
sys.exit(1 if FAILED else 0)
